"""Populate a real Word TOC from heading styles and verified PDF pagination.
Word can rebuild the TOC field when headings change. The delivered cached field
is generated automatically and checked against the actual rendered pages.
"""
from pathlib import Path
import json,re,subprocess,os
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER,WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
ROOT=Path(__file__).resolve().parents[1]
NORM=lambda s:re.sub(r'[^а-яёa-z0-9]','',s.lower())

def element(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e
def run(parent,text=None,node=None):
 r=element('r');parent.append(r)
 if text is not None:
  t=element('t');t.text=text;r.append(t)
 if node is not None:r.append(node)
 return r
def heads(doc):return [(p.text,int(p.style.name[-1])) for p in doc.paragraphs if p.style.name in ['Heading 1','Heading 2']]
def cache(doc,entries,pages):
 body=doc._element.body
 existing=body.xpath('./w:sdt[w:sdtPr/w:tag[@w:val="UITU_TOC"]]')
 if existing:old=existing[0]
 else:old=next(p._p for p in doc.paragraphs if any('TOC ' in (x.text or '') for x in p._p.xpath('.//w:instrText')))
 index=body.index(old);body.remove(old)
 for depth in [1,2]:
  name=f'TOC {depth}'
  style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
  style.font.name='Times New Roman';style.font.size=Pt(14);style.font.color.rgb=RGBColor(0,0,0)
  f=style.paragraph_format;f.first_line_indent=Cm(0);f.left_indent=Cm(0 if depth==1 else .5);f.line_spacing=1;f.space_before=Pt(0);f.space_after=Pt(0);f.keep_with_next=False
  f.tab_stops.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 bookmarks=[]
 for i,p in enumerate([p for p in doc.paragraphs if p.style.name in ['Heading 1','Heading 2']],1):
  for b in p._p.xpath('./w:bookmarkStart|./w:bookmarkEnd'):p._p.remove(b)
  p._p.insert(1,element('bookmarkStart',id=100+i,name=f'_TocUITU{i}'));p._p.append(element('bookmarkEnd',id=100+i))
 sdt=element('sdt');pr=element('sdtPr');pr.append(element('tag',val='UITU_TOC'));obj=element('docPartObj');obj.append(element('docPartGallery',val='Table of Contents'));pr.append(obj);sdt.append(pr)
 content=element('sdtContent');sdt.append(content)
 for i,(title,depth) in enumerate(entries):
  p=element('p');ppr=element('pPr');ppr.append(element('pStyle',val=doc.styles[f'TOC {depth}'].style_id));p.append(ppr)
  if i==0:
   run(p,node=element('fldChar',fldCharType='begin'))
   ins=element('instrText');ins.set(qn('xml:space'),'preserve');ins.text=' TOC \\o "1-2" \\h \\z \\u '
   run(p,node=ins);run(p,node=element('fldChar',fldCharType='separate'))
  link=element('hyperlink',anchor=f'_TocUITU{i+1}');p.append(link);run(link,title)
  run(p,node=element('tab'));run(p,str(pages.get(title,1)))
  if i==len(entries)-1:run(p,node=element('fldChar',fldCharType='end'))
  content.append(p)
 body.insert(index,sdt)
 update=doc.settings.element.find(qn('w:updateFields'))
 if update is None:doc.settings.element.append(element('updateFields',val='true'))

def get_pages(pdf,entries):
 found={}
 with fitz.open(pdf) as d:
  for title,depth in entries:
   candidates=[];target=NORM(title)
   for n,page in enumerate(d,1):
    lines=page.get_text().splitlines()
    for i in range(len(lines)):
     for k in range(1,4):
      if NORM(''.join(lines[i:i+k]))==target:candidates.append(n)
   # Cached TOC entries contain a page number and do not match the heading.
   if not candidates:raise AssertionError(('Heading not found in PDF',title))
   found[title]=max(candidates)
  return len(d),found

def finalize(path):
 doc=Document(path);entries=heads(doc);pages={};count=1
 for iteration in range(6):
  cache(doc,entries,pages)
  for p in doc.paragraphs:
   if p.text.startswith('Объём отчёта:'):
    p.text=re.sub(r'Объём отчёта: .*? с\.',f'Объём отчёта: {count} с.',p.text)
  doc.save(path)
  env={**os.environ,'FONTCONFIG_FILE':'/workspace/.uitu-tools/fonts.conf','XDG_CACHE_HOME':'/workspace/.uitu-tools/cache'}
  result=subprocess.run(['soffice',f'-env:UserInstallation=file:///tmp/uitu-finalize-{path.stem}','--headless','--convert-to','pdf','--outdir',str(path.parent),str(path)],env=env,text=True,capture_output=True,timeout=120)
  if result.returncode:raise RuntimeError(result.stderr)
  new_count,new_pages=get_pages(path.with_suffix('.pdf'),entries)
  if new_count==count and new_pages==pages:
   print(path.name,new_count,'pages; TOC verified')
   return {'file':path.name,'pages':new_count,'toc_entries':len(entries),'toc_pages':new_pages,'iterations':iteration+1}
  count,pages=new_count,new_pages
 raise RuntimeError('Pagination did not stabilize: '+str(path))

if __name__=='__main__':
 results=[finalize(path) for path in sorted((ROOT/'reports').glob('*.docx'))]
 (ROOT/'evidence/render-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
