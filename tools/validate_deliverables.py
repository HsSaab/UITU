"""Validate delivered artifacts against the rendered documents and live evidence."""
from pathlib import Path
import re,json,hashlib
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import fitz
from PIL import Image
from finalize_reports import NORM
ROOT=Path(__file__).resolve().parents[1]
NAMES=['Тарасов Арсений Максимович','Дементиевский Дмитрий Максимович','Тыцкий Богдан Максимович']
results=[]
for number in [4,5,6]:
 path=ROOT/'reports'/f'УИТУ_Практическая_{number}_ТехноСервис.docx'
 with ZipFile(path) as z:assert z.testzip() is None
 doc=Document(path);pdf=fitz.open(path.with_suffix('.pdf'))
 text='\n'.join(p.text for p in doc.paragraphs)
 xml_text=''.join(doc._element.xpath('.//w:t/text()'))
 assert not re.search(r'[\u2013\u2014\[\]]',xml_text),(path.name,'forbidden punctuation in DOCX')
 assert not re.search(r'PAGE_COUNT|FIGURE_COUNT|TABLE_COUNT|SOURCE_COUNT|Содержание обновляется',text)
 assert all(n in xml_text for n in NAMES)
 assert 'ИКБО-74-23' in xml_text and 'Демидов Н.А.' in xml_text
 assert 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ' not in xml_text
 for title in ['ЦЕЛЬ ПРАКТИЧЕСКОЙ РАБОТЫ','ВЫПОЛНЕНИЕ РАБОТЫ','Вывод по работе']:
  assert any(p.text==title and p.style.name=='Heading 1' for p in doc.paragraphs)
 sec=doc.sections[0]
 for attr,value in [('page_width',21),('page_height',29.7),('left_margin',3),('right_margin',1.5),('top_margin',2),('bottom_margin',2)]:
  assert abs(getattr(sec,attr).cm-value)<.01,(path.name,attr)
 assert sec.different_first_page_header_footer
 normal=doc.styles['Normal'];assert normal.font.name=='Times New Roman' and normal.font.size.pt==14
 assert normal.paragraph_format.line_spacing==1.5
 assert abs(normal.paragraph_format.first_line_indent.cm-1.25)<.01
 figure_paragraphs=[p for p in doc.paragraphs if re.match(r'^Рисунок \d+ -',p.text)]
 figures=[p.text for p in figure_paragraphs]
 assert doc.styles['Caption'].font.italic is True
 for caption in figure_paragraphs:
  assert caption.style.name=='Caption'
  assert all(run.italic is not False for run in caption.runs)
 tables=[p.text for p in doc.paragraphs if p.text.startswith('Таблица ')]
 section=None;section_counts={};table_numbers=[]
 for paragraph in doc.paragraphs:
  if paragraph.style.name in ['Heading 1','Heading 2','Heading 3']:
   match=re.match(r'^(\d+)(?:\.|\s)',paragraph.text)
   if match:section=int(match.group(1))
  if paragraph.text.startswith('Таблица '):
   match=re.match(r'^Таблица (\d+)\.(\d+) - ',paragraph.text)
   assert match,(path.name,'table must use section.table numbering',paragraph.text)
   assert int(match.group(1))==section,(path.name,'wrong table section',paragraph.text)
   section_counts[section]=section_counts.get(section,0)+1
   assert int(match.group(2))==section_counts[section]
   table_numbers.append('.'.join(match.groups()))
 sources=[p for p in doc.paragraphs if p.style.name=='Bibliography']
 assert not sources
 assert len(figures)+1==len(doc.inline_shapes)
 assert doc.inline_shapes[0]._inline.docPr.get('descr')=='Эмблема РТУ МИРЭА'
 for i in range(1,len(figures)+1):assert re.search(r'рисунке '+str(i)+r'\b',text)
 for label in table_numbers:
  assert f'в таблице {label}.' in text
  assert any(page.search_for('Таблица '+label+' -') for page in pdf),(path.name,label)
 # Confirm italics in the actual exported PDF, including wrapped captions.
 for caption in figures:
  target=NORM(caption);found=False
  for page in pdf:
   lines=[line for block in page.get_text('dict')['blocks'] if block['type']==0 for line in block['lines']]
   for start,line in enumerate(lines):
    first=''.join(span['text'] for span in line['spans'])
    if not first.startswith('Рисунок '):continue
    for count in range(1,5):
     selected=lines[start:start+count]
     spans=[span for item in selected for span in item['spans'] if span['text'].strip()]
     if NORM(''.join(span['text'] for span in spans))==target:
      assert all(span['flags'] & 2 for span in spans),(path.name,'nonitalic PDF caption',caption)
      found=True;break
   if found:break
  assert found,(path.name,'caption not found in PDF',caption)
 headings=[p for p in doc.paragraphs if p.style.name in ['Heading 1','Heading 2','Heading 3']]
 cached=doc._element.body.xpath('./w:sdt[w:sdtPr/w:tag[@w:val="UITU_TOC"]]/w:sdtContent/w:p')
 assert len(headings)==len(cached)
 assert any('TOC ' in (x.text or '') for x in doc._element.body.xpath('.//w:instrText'))
 from finalize_reports import get_pages
 _,pages=get_pages(path.with_suffix('.pdf'),[(p.text,1) for p in headings])
 for p,entry in zip(headings,cached):
  links=entry.xpath('./w:hyperlink');assert len(links)==1
  link=links[0]
  title=link.xpath('./w:r/w:t')[0].text
  number_text=link.xpath('./w:r/w:t')[-1].text
  assert title==p.text and int(number_text)==pages[p.text]
  assert link.xpath('./w:r/w:tab') and not entry.xpath('./w:r/w:t')
  anchor=link.get(qn('w:anchor'))
  starts=p._p.xpath('./w:bookmarkStart');ends=p._p.xpath('./w:bookmarkEnd')
  assert any(s.get(qn('w:name'))==anchor and any(e.get(qn('w:id'))==s.get(qn('w:id')) for e in ends) for s in starts)
 # Check the actual PDF link hit areas, including page numbers and leaders.
 toc_start=next(i for i,page in enumerate(pdf) if any(line.strip()=='СОДЕРЖАНИЕ' for line in page.get_text().splitlines()))
 toc_end=pages['ЦЕЛЬ ПРАКТИЧЕСКОЙ РАБОТЫ']-1
 toc_lines=[]
 for page_index in range(toc_start,toc_end):
  for block in pdf[page_index].get_text('rawdict')['blocks']:
   if block['type']!=0:continue
   for line in block['lines']:
    chars=[c for span in line['spans'] for c in span['chars']]
    toc_lines.append((page_index,chars,''.join(c['c'] for c in chars)))
 from finalize_reports import NORM
 for heading in headings:
  matches=[row for row in toc_lines if NORM(row[2]).startswith(NORM(heading.text))]
  assert len(matches)==1,(path.name,heading.text,'TOC row not found')
  page_index,chars,line=matches[0]
  digits=re.search(r'(\d+)\s*$',line);assert digits
  assert int(digits.group(1))==pages[heading.text]
  dots=list(re.finditer(r'\.{3,}',line));assert dots
  indices=[0,digits.start(),(dots[-1].start()+dots[-1].end())//2]
  for index in indices:
   rect=fitz.Rect(chars[index]['bbox']);point=(rect.tl+rect.br)/2
   active=[l for l in pdf[page_index].get_links() if l['kind']==fitz.LINK_GOTO and point in l['from']]
   assert active and all(l['page']==pages[heading.text]-1 for l in active),(path.name,heading.text,'nonclickable or incorrect TOC target',index)
 for i,page in enumerate(pdf):
  assert 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ' not in page.get_text()
  assert not re.search(r'[\u2013\u2014\[\]]',page.get_text()),(path.name,i+1,'forbidden punctuation in PDF')
  assert len(page.get_text().strip())>220,(path.name,'unexpected short page',i+1)
  bottom=page.get_text(clip=fitz.Rect(0,795,page.rect.width,page.rect.height)).strip()
  assert bottom==('' if i==0 else str(i+1)),(path.name,i+1,bottom)
  for block in page.get_text('dict')['blocks']:
   if block['type']==0:
    x0,y0,x1,y1=block['bbox'];assert x0>=25 and x1<=page.rect.width-20 and y0>=20 and y1<=page.rect.height-10
 hashes={suffix:hashlib.sha256(path.with_suffix(suffix).read_bytes()).hexdigest() for suffix in ['.docx','.pdf']}
 results.append({'report':number,'pages':len(pdf),'figures':len(figures),'figure_captions':'italic in DOCX and PDF','tables':len(tables),'table_numbers':table_numbers,'sources':len(sources),'toc':'verified against PDF; title, leaders and page number clickable','toc_links':len(headings),'sha256':hashes})
 pdf.close()
index=json.loads((ROOT/'diagrams/idef0-index.json').read_text());assert len(index)==12
for key in ['vps','monitoring']:
 nodes={m['node'] for m in index if m['stem'].startswith(key)};assert nodes=={'A-0','A0','A1','A2','A3','A4'}
for item in index:
 for extension in ['.drawio','.svg']:etree.parse(str(ROOT/'diagrams'/(item['stem']+extension)))
 tree=etree.parse(str(ROOT/'diagrams'/(item['stem']+'.drawio')))
 cells={c.get('id'):c for c in tree.iter('mxCell')}
 for cell in cells.values():
  if cell.get('edge')!='1':continue
  for side in ['source','target']:
   key=cell.get(side)
   if not key:continue
   box=cells[key].find('mxGeometry')
   x,y,w,h=[float(box.get(a)) for a in ['x','y','width','height']]
   point=cell.find(f'mxGeometry/mxPoint[@as="{side}Point"]')
   px,py=float(point.get('x')),float(point.get('y'))
   assert x<=px<=x+w and y<=py<=y+h,(item['stem'],'arrow outside function')
   assert min(abs(px-x),abs(px-x-w),abs(py-y),abs(py-y-h))<.01
 with Image.open(ROOT/'diagrams/png'/(item['stem']+'.png')) as im:assert im.width>=2400
api=json.loads((ROOT/'evidence/api-results.json').read_text());assert api['executed']==20 and api['failed']==api['errors']==0
browser=json.loads((ROOT/'evidence/browser-results.json').read_text());assert browser['executed']==13 and browser['failed']==0
units=json.loads((ROOT/'evidence/unit-results.json').read_text());assert len(units)==10 and all(x['status']=='passed' for x in units)
for p in (ROOT/'docs/practices').glob('*.docx'):
 originals=list(Path('/workspace/attachments').rglob(p.name))
 if originals:assert p.read_bytes()==originals[0].read_bytes()
(ROOT/'evidence/document-validation.json').write_text(json.dumps({'reports':results,'idef0_diagrams':12,'unit_tests':10,'api_tests':20,'browser_checks':13,'status':'passed'},ensure_ascii=False,indent=2))
for r in results:print(f'Report {r["report"]}: {r["pages"]} pages, {r["figures"]} figures, {r["tables"]} tables; validated')
print('All documents, TOCs, page numbers, diagrams and test evidence validated.')
