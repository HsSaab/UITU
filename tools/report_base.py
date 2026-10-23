from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from text_cleanup import clean_text

ROOT=Path(__file__).resolve().parents[1]
NAMES=['Тарасов Арсений Максимович','Дементиевский Дмитрий Максимович','Тыцкий Богдан Максимович']

def font(style,size=14):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    style.element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        style.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)

class Report:
    def __init__(self,num,title,keywords,abstract):
        self.num=num;self.title=title;self.d=Document();self.figures=0;self.tables=0
        self.current_section=None;self.section_tables={};self.table_numbers=[]
        self.section_figures={};self.figure_numbers=[]
        sec=self.d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.left_margin=Cm(3);sec.right_margin=Cm(1.5);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
        sec.footer_distance=Cm(1);sec.different_first_page_header_footer=True
        for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption','TOC Heading']:
            s=self.d.styles[name];font(s)
            f=s.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1.5
            f.first_line_indent=Cm(1.25);f.widow_control=True
        normal=self.d.styles['Normal'];normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        for name in ['Heading 1','Heading 2','Heading 3']:
            s=self.d.styles[name];s.font.bold=True
            s.paragraph_format.keep_with_next=True;s.paragraph_format.line_spacing=1
        self.d.styles['Heading 1'].paragraph_format.first_line_indent=Cm(1.25)
        self.d.styles['Heading 1'].paragraph_format.page_break_before=True
        self.d.styles['Caption'].paragraph_format.first_line_indent=Cm(0)
        self.d.styles['Caption'].paragraph_format.left_indent=Cm(0)
        self.d.styles['Caption'].paragraph_format.right_indent=Cm(0)
        self.d.styles['Caption'].paragraph_format.line_spacing=1
        self.d.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.d.styles['Caption'].font.italic=True
        self.d.styles['Caption'].font.bold=False
        footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Cm(0)
        r=footer.add_run();r.font.name='Times New Roman';r.font.size=Pt(14)
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
        # Russian language enables correct hyphenation/spell checking in Word.
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');normal.element.get_or_add_rPr().append(lang)
        self.titlepage()
        self.heading('СОДЕРЖАНИЕ',center=True,style='TOC Heading')
        # A genuine Word TOC field, updated by LibreOffice before delivery.
        p=self.d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
        begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
        instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-3" \\h \\z \\u '
        sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
        end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
        for node in [begin,instr,sep]:p.add_run()._r.append(node)
        p.add_run('Содержание обновляется по заголовкам документа.')
        p.add_run()._r.append(end)
        p.paragraph_format.page_break_before=False
        # The TOC heading is a standalone structural heading, outside the TOC itself.
        p0=p._p.getprevious();p0.get_or_add_pPr().append(self._bool('pageBreakBefore'))
        self.d.core_properties.author='; '.join(NAMES);self.d.core_properties.title=clean_text(f'Практическая работа №{num}. {title}')
        self.d.core_properties.subject='Управление информационно-технологическими услугами. ООО «ТехноСервис»'
    def _bool(self,name):return OxmlElement('w:'+name)
    def titlepage(self):
        p=self.p('',align='center');p.paragraph_format.line_spacing=1
        picture=p.add_run().add_picture(str(ROOT/'assets/mirea-emblem.png'),height=Cm(2.4))
        picture._inline.docPr.set('descr','Эмблема РТУ МИРЭА')
        self.p('МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ',align='center')
        self.p('Федеральное государственное бюджетное образовательное учреждение высшего образования',align='center')
        self.p('«МИРЭА - Российский технологический университет»',align='center',bold=True)
        p=self.p('РТУ МИРЭА',align='center',bold=True)
        borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
        for key,value in [('val','double'),('sz','6'),('space','8'),('color','000000')]:bottom.set(qn('w:'+key),value)
        borders.append(bottom);p._p.get_or_add_pPr().append(borders)
        self.p('')
        self.p('Институт информационных технологий (ИТ)',align='center')
        self.p('Кафедра корпоративных информационных систем (КИС)',align='center')
        self.p(f'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ №{self.num}',align='center',bold=True)
        self.p('по дисциплине',align='center',bold=True)
        self.p('«Управление информационно-технологическими услугами»',align='center',bold=True)
        self.p('')
        table=self.d.add_table(rows=2,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        for col,width in zip(table.columns,[6.3,10.2]):col.width=Cm(width)
        values=[['Выполнили студенты\nгруппы ИКБО-74-23','\n'.join(NAMES)],['Принял преподаватель','ассистент Демидов Н.А.']]
        for i,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(self._bool('cantSplit'))
            for cell,width,value in zip(row.cells,[6.3,10.2],values[i]):
                cell.width=Cm(width);cell.text=value
                for p in cell.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.2
                    p.paragraph_format.space_before=Pt(15 if i else 0);p.paragraph_format.space_after=Pt(0)
                    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(14)
        p=self.p('Москва 2026',align='center');p.paragraph_format.space_before=Cm(1.2)
    def p(self,text='',align=None,bold=False,size=14,keep=False):
        p=self.d.add_paragraph();r=p.add_run(clean_text(text));r.bold=bold;r.font.size=Pt(size)
        if align:p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT}[align];p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.keep_with_next=keep
        return p
    def heading(self,text,center=False,style='Heading 1'):
        section=re.match(r'^(\d+)(?:\.|\s)',text)
        if section and style in ['Heading 1','Heading 2','Heading 3']:
            self.current_section=int(section.group(1))
        p=self.d.add_paragraph(clean_text(text),style=style)
        if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.space_after=Pt(14)
        return p
    def sub(self,text,newpage=False):
        p=self.heading(text,style='Heading 2')
        if newpage:p.paragraph_format.page_break_before=True
        return p
    def minor(self,text,newpage=False):
        p=self.heading(text,style='Heading 3')
        if newpage:p.paragraph_format.page_break_before=True
        return p
    def bullet(self,text):
        return self.p('- '+text)
    def table(self,caption,headers,rows,widths=None,size=11.5):
        if self.current_section is None:
            raise ValueError('Таблица должна находиться в нумерованном разделе: '+caption)
        self.tables+=1
        self.section_tables[self.current_section]=self.section_tables.get(self.current_section,0)+1
        number=f'{self.current_section}.{self.section_tables[self.current_section]}'
        self.table_numbers.append(number)
        self.p(f'Данные представлены в таблице {number}.',keep=True)
        cap=self.p(f'Таблица {number} - {caption}',align='left',keep=True);cap.paragraph_format.line_spacing=1
        table=self.d.add_table(rows=1,cols=len(headers));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        if widths is None:widths=[16.5/len(headers)]*len(headers)
        for c,w in zip(table.columns,widths):c.width=Cm(w)
        for cell,h in zip(table.rows[0].cells,headers):cell.text=clean_text(h)
        for values in rows:
            for cell,v in zip(table.add_row().cells,values):cell.text=clean_text(v)
        trpr=table.rows[0]._tr.get_or_add_trPr();trpr.append(self._bool('tblHeader'))
        for ri,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(self._bool('cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    p.paragraph_format.keep_with_next=(ri<len(table.rows)-1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:r.font.size=Pt(size);r.font.name='Times New Roman';r.bold=(ri==0)
        return table
    def figure(self,path,caption,width=16.5):
        from PIL import Image
        if self.current_section is None:
            raise ValueError('Иллюстрация должна находиться в нумерованном разделе: '+caption)
        self.figures+=1
        self.section_figures[self.current_section]=self.section_figures.get(self.current_section,0)+1
        number=f'{self.current_section}.{self.section_figures[self.current_section]}'
        self.figure_numbers.append(number)
        self.p(f'Результат представлен на рисунке {number}.',keep=True)
        self.p('',align='center',keep=True)
        p=self.d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0)
        p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        p.paragraph_format.line_spacing=1
        with Image.open(path) as im:w,h=im.size
        actual=min(width,19*w/h)
        picture=p.add_run().add_picture(str(path),width=Cm(actual))
        # Set every wrap/effect distance explicitly. LibreOffice otherwise adds
        # default graphic margins that shift a full-width inline picture.
        for name in ('distT','distB','distL','distR'):picture._inline.set(name,'0')
        extent=OxmlElement('wp:effectExtent')
        for name in ('l','t','r','b'):extent.set(name,'0')
        picture._inline.insert(1,extent)
        picture._inline.docPr.set('descr',f'Рисунок {number}. '+clean_text(caption))
        # Native Word picture outline: 0.25 pt is the minimum standard line weight.
        outline=OxmlElement('a:ln');outline.set('w','3175')
        fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr')
        color.set('val','000000');fill.append(color);outline.append(fill)
        dash=OxmlElement('a:prstDash');dash.set('val','solid');outline.append(dash)
        picture._inline.xpath('.//pic:spPr')[0].append(outline)
        # The requested caption separator is an em dash; body text keeps its
        # earlier punctuation cleanup through clean_text.
        label=self.d.add_paragraph(f'Рисунок {number} — '+clean_text(caption),style='Caption')
        label.paragraph_format.keep_with_next=True;label.paragraph_format.keep_together=True
        self.p('',align='center')
    def save(self):
        path=ROOT/'reports'/f'УИТУ_Практическая_{self.num}_ТехноСервис.docx'
        self.d.save(path);return path
